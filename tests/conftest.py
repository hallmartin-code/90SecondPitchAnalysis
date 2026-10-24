from __future__ import annotations

import io
import json
from pathlib import Path

import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pitch90.config import ASSETS_DIR, SAMPLES_DIR, Settings
from pitch90.models import LLMAnalysis

FONT = ASSETS_DIR / "fonts" / "Carlito-Regular.ttf"


@pytest.fixture
def settings() -> Settings:
    return Settings(ANTHROPIC_API_KEY=None, PITCH90_SOFFICE_PATH="", _env_file=None)


def text_image(text: str, size=(1400, 500)) -> bytes:
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((40, 180), text, font=ImageFont.truetype(str(FONT), 64), fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def pdf_bytes() -> bytes:
    """Page 1-2 text, page 3 image-only (text in picture), page 4 blank."""
    doc = pymupdf.open()
    for marker in ("Alpha page talks about the customer problem in detail",
                   "Beta page shows revenue of $2.4M in FY2025 and 14 customers"):
        page = doc.new_page()
        page.insert_text((72, 100), marker, fontsize=14)
    page = doc.new_page()
    page.insert_image(page.rect, stream=text_image("Pilot cut downtime 27%"))
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def pptx_bytes() -> bytes:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Slide one title"
    s1.placeholders[1].text = "Visible body text on slide one"
    s1.notes_slide.notes_text_frame.text = "Presenter notes for slide one"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Comparison"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, row in enumerate([("Metric", "Value"), ("ARR", "$3.1M")]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Growth"
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Revenue ($M)", (1.0, 2.5))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)

    s4 = prs.slides.add_slide(prs.slide_layouts[1])
    s4.shapes.title.text = "Hidden appendix"
    s4.placeholders[1].text = "Not shown when presenting"
    s4._element.set("show", "0")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def demo_json() -> dict:
    return json.loads((SAMPLES_DIR / "demo_analysis.json").read_text(encoding="utf-8"))


@pytest.fixture
def demo_llm(demo_json) -> LLMAnalysis:
    return LLMAnalysis.model_validate(demo_json)


@pytest.fixture(scope="session")
def sample_deck():
    from pitch90.ingest import extract_deck

    path: Path = SAMPLES_DIR / "tidewell_sample_deck.pptx"
    s = Settings(ANTHROPIC_API_KEY=None, _env_file=None)
    return extract_deck(path.name, path.read_bytes(), s)
