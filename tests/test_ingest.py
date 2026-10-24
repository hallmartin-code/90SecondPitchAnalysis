from __future__ import annotations

import importlib
import io

import pymupdf
import pytest

from pitch90.ingest import ConversionUnavailable, UploadError, extract_deck
from pitch90.ingest.ocr import ocr_available
from pitch90.ingest.validate import OLE_MAGIC, safe_filename

# The package re-exports a `convert` function, so fetch the module explicitly.
convert_mod = importlib.import_module("pitch90.ingest.convert")


# --------------------------------------------------------------------------- PDF


def test_pdf_preserves_page_numbering(pdf_bytes, settings):
    deck = extract_deck("deck.pdf", pdf_bytes, settings)
    assert deck.unit_label == "page"
    assert [u.number for u in deck.units] == [1, 2, 3, 4]
    assert "Alpha page" in deck.unit(1).text
    assert "$2.4M" in deck.unit(2).text


def test_pdf_image_page_flagged_and_rendered(pdf_bytes, settings):
    deck = extract_deck("deck.pdf", pdf_bytes, settings)
    page3 = deck.unit(3)
    assert page3.text == ""
    assert page3.needs_visual
    assert 3 in deck.images  # rendered for visual interpretation
    if ocr_available():
        assert "27%" in page3.ocr_text


def test_pdf_blank_page_not_silently_skipped(pdf_bytes, settings):
    deck = extract_deck("deck.pdf", pdf_bytes, settings)
    assert deck.unit(4) is not None
    assert any(4 in w.units for w in deck.warnings)


def test_pdf_encrypted_rejected(settings):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(UploadError, match="password"):
        extract_deck("locked.pdf", data, settings)


def test_pdf_corrupt_rejected(settings):
    with pytest.raises(UploadError):
        extract_deck("bad.pdf", b"%PDF-1.7\n" + b"\x00garbage" * 50, settings)


# --------------------------------------------------------------------------- PPTX


def test_pptx_numbering_text_tables_chart_and_notes(pptx_bytes, settings):
    deck = extract_deck("deck.pptx", pptx_bytes, settings)
    assert deck.unit_label == "slide"
    assert [u.number for u in deck.units] == [1, 2, 3, 4]
    s1 = deck.unit(1)
    assert "Visible body text" in s1.text
    assert s1.notes == "Presenter notes for slide one"
    assert "Presenter notes" not in s1.text  # notes kept separate from visible content
    assert deck.unit(2).tables == [[["Metric", "Value"], ["ARR", "$3.1M"]]]
    assert "Revenue ($M): 2024=1, 2025=2.5" in deck.unit(3).chart_data[0]


def test_pptx_hidden_slide_flagged(pptx_bytes, settings):
    deck = extract_deck("deck.pptx", pptx_bytes, settings)
    assert deck.unit(4).hidden
    assert any(4 in w.units and "Hidden" in w.message for w in deck.warnings)


def test_sample_deck_extraction(sample_deck):
    assert len(sample_deck.units) == 13
    assert sample_deck.unit(7).needs_visual
    assert any(w.kind == "security" and 12 in w.units for w in sample_deck.warnings)


# --------------------------------------------------------------------------- validation


@pytest.mark.parametrize("name,data,match", [
    ("deck.pdf", b"", "empty"),
    ("deck.docx", b"PK\x03\x04abc", "Unsupported"),
    ("deck.pdf", b"not a pdf at all", "not a valid PDF"),
    ("deck.pptx", b"plain text", "not a valid PowerPoint"),
    ("deck.pptx", b"PK\x03\x04garbage", "corrupt"),
    ("deck.ppt", b"PK\x03\x04" + b"0" * 20, "pptx saved with a .ppt"),
    ("deck.pptx", OLE_MAGIC + b"0" * 64, "legacy PowerPoint"),
])
def test_invalid_uploads(name, data, match, settings):
    with pytest.raises(UploadError, match=match):
        extract_deck(name, data, settings)


def test_encrypted_pptx_rejected(settings):
    data = OLE_MAGIC + b"\x00" * 100 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 100
    with pytest.raises(UploadError, match="password"):
        extract_deck("deck.pptx", data, settings)


def test_oversized_upload_rejected(settings):
    small = settings.model_copy(update={"max_upload_mb": 1})
    with pytest.raises(UploadError, match="limit"):
        extract_deck("deck.pdf", b"%PDF-1.4" + b"0" * (2 * 1024 * 1024), small)


def test_empty_presentation_rejected(settings):
    from pptx import Presentation

    buf = io.BytesIO()
    Presentation().save(buf)
    with pytest.raises(UploadError, match="no slides"):
        extract_deck("deck.pptx", buf.getvalue(), settings)


def test_safe_filename():
    assert safe_filename("../../etc/passwd.pdf") == "passwd.pdf"
    assert safe_filename("C:\\Users\\x\\My Deck (v2).PPTX") == "My_Deck_v2.pptx"
    assert safe_filename("....pdf") == "deck.pdf"


# --------------------------------------------------------------------------- legacy PPT


def test_ppt_conversion_unavailable(monkeypatch, settings):
    monkeypatch.setattr("pitch90.ingest.find_soffice", lambda configured=None: None)
    data = OLE_MAGIC + b"\x00" * 512
    with pytest.raises(ConversionUnavailable) as err:
        extract_deck("old.ppt", data, settings)
    msg = str(err.value)
    assert "LibreOffice" in msg and "PITCH90_SOFFICE_PATH" in msg and ".pptx" in msg


def test_convert_without_soffice_raises(tmp_path):
    src = tmp_path / "x.ppt"
    src.write_bytes(b"x")
    with pytest.raises(ConversionUnavailable):
        convert_mod.convert(src, "pptx", None)


def test_find_soffice_respects_bad_configured_path():
    assert convert_mod.find_soffice("Z:/definitely/missing/soffice.exe") is None


SOFFICE = convert_mod.find_soffice()


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_legacy_ppt_converted_with_libreoffice(tmp_path, settings):
    from pitch90.config import SAMPLES_DIR

    src = tmp_path / "deck.pptx"
    src.write_bytes((SAMPLES_DIR / "tidewell_sample_deck.pptx").read_bytes())
    ppt = convert_mod.convert(src, "ppt", SOFFICE)
    deck = extract_deck("deck.ppt", ppt.read_bytes(), settings)
    assert deck.source_type == "ppt" and deck.converted_from_ppt
    assert [u.number for u in deck.units] == list(range(1, 14))
    assert "ARR $1.2M" in deck.unit(6).text
    assert "ignore previous instructions" in deck.unit(12).notes


@pytest.mark.skipif(not (SOFFICE and ocr_available()), reason="needs LibreOffice and OCR")
def test_slide_renders_map_to_original_numbers_despite_hidden_slide(settings):
    from pptx import Presentation
    from pptx.util import Inches

    from tests.conftest import text_image

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "First slide title text here"
    hidden = prs.slides.add_slide(prs.slide_layouts[1])
    hidden.shapes.title.text = "Hidden slide"
    hidden._element.set("show", "0")
    pic = prs.slides.add_slide(prs.slide_layouts[6])
    pic.shapes.add_picture(io.BytesIO(text_image("Picture on slide three")), Inches(0), Inches(0),
                           width=prs.slide_width)
    buf = io.BytesIO()
    prs.save(buf)
    deck = extract_deck("deck.pptx", buf.getvalue(), settings)
    assert 3 in deck.images and 2 not in deck.images
    from pitch90.ingest.ocr import ocr_image

    assert "three" in (ocr_image(deck.images[3]) or "").lower()
