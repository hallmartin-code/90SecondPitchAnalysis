"""PPTX extraction with python-pptx: text, tables, chart data, pictures and notes."""

from __future__ import annotations

import io
import logging
import re
import time

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from pitch90.ingest.ocr import ocr_image
from pitch90.ingest.validate import UploadError
from pitch90.models import ExtractedDeck, ExtractionWarning, UnitContent

log = logging.getLogger(__name__)

MIN_TEXT_CHARS = 40
IMAGE_COVERAGE_VISUAL = 0.30


def _alnum(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9]", text))


def _shape_type(shape):
    """python-pptx raises NotImplementedError for some unusual shapes; treat those as unknown."""
    try:
        return shape.shape_type
    except Exception:
        return None


def _iter_shapes(shapes):
    for shape in shapes:
        if _shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _iter_shapes(shape.shapes)
            except Exception:
                continue
        else:
            yield shape


def _chart_summary(chart) -> str:
    parts = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            parts.append(f"Chart: {chart.chart_title.text_frame.text.strip()}")
        else:
            parts.append("Chart")
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        for series in plot.series:
            values = list(series.values)
            pairs = ", ".join(
                f"{cat}={val:g}" if isinstance(val, (int, float)) else f"{cat}={val}"
                for cat, val in zip(categories, values)
            )
            parts.append(f"{series.name}: {pairs}")
    except Exception:
        parts.append("(chart data could not be read)")
    return "; ".join(parts)


def _to_jpeg(blob: bytes, max_px: int = 1400) -> bytes | None:
    try:
        img = Image.open(io.BytesIO(blob))
        img.thumbnail((max_px, max_px))
        out = io.BytesIO()
        img.convert("RGB").save(out, format="JPEG", quality=80)
        return out.getvalue()
    except Exception:
        return None


def _extract_slide(slide, number: int, slide_area: int, allow_ocr: bool
                   ) -> tuple[UnitContent, bytes | None, bool]:
    """Extract one slide. Returns (unit, largest picture as JPEG or None, OCR-was-unavailable)."""
    unit = UnitContent(number=number, hidden=slide._element.get("show") == "0")
    if unit.hidden:
        unit.warnings.append("Hidden slide (included in extraction).")
    texts: list[tuple[int, int, str]] = []
    picture_blobs: list[tuple[int, bytes]] = []
    picture_area = 0
    unreadable_graphics = 0
    skipped_shapes = 0

    for shape in _iter_shapes(slide.shapes):
        try:
            shape_type = _shape_type(shape)
            top, left = int(shape.top or 0), int(shape.left or 0)
            area = int(shape.width or 0) * int(shape.height or 0)
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if text:
                    texts.append((top, left, text))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                unit.tables.append(rows)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                unit.chart_data.append(_chart_summary(shape.chart))
            elif shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    picture_blobs.append((area, shape.image.blob))
                    picture_area += area
                except Exception:
                    pass
            elif type(shape).__name__ == "GraphicFrame" or shape_type in (
                MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
            ):
                # SmartArt diagrams and OLE objects: python-pptx cannot read their text.
                unreadable_graphics += 1
        except Exception:
            skipped_shapes += 1

    texts.sort(key=lambda t: (t[0], t[1]))
    unit.text = "\n".join(t[2] for t in texts)
    unit.image_count = len(picture_blobs)
    try:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            unit.notes = slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        unit.warnings.append("Speaker notes could not be read.")
    if unreadable_graphics:
        unit.warnings.append(
            f"{unreadable_graphics} diagram/SmartArt/embedded object(s) whose text could not be extracted."
        )
    if skipped_shapes:
        unit.warnings.append(f"{skipped_shapes} shape(s) could not be read and were skipped.")

    picture_blobs.sort(key=lambda p: p[0], reverse=True)
    picture = _to_jpeg(picture_blobs[0][1]) if picture_blobs else None

    no_ocr = False
    visible_chars = _alnum(unit.text) + sum(_alnum(" ".join(r)) for t in unit.tables for r in t)
    low_text = visible_chars < MIN_TEXT_CHARS and not unit.chart_data
    if low_text and picture_blobs and allow_ocr:
        ocr_parts = []
        for _, blob in picture_blobs[:3]:
            result = ocr_image(blob)
            if result is None:
                no_ocr = True
                break
            if result:
                ocr_parts.append(result)
        unit.ocr_text = "\n".join(ocr_parts)
        if unit.ocr_text:
            unit.warnings.append("Text recovered from slide images by OCR; check for recognition errors.")
    if low_text and not unit.ocr_text:
        unit.warnings.append("Little or no readable text on this slide.")

    reasons = []
    if low_text:
        reasons.append("little extractable text")
    if picture_area / slide_area >= IMAGE_COVERAGE_VISUAL:
        reasons.append(f"images cover {picture_area / slide_area:.0%} of the slide")
    if unreadable_graphics:
        reasons.append("diagram or embedded object")
    if reasons:
        unit.needs_visual = True
        unit.visual_reason = "; ".join(reasons)
    return unit, picture, no_ocr


def extract_pptx(
    data: bytes, filename: str, max_pages: int, source_type: str = "pptx", budget_seconds: float = 120.0,
) -> tuple[ExtractedDeck, dict[int, bytes]]:
    """Return the extracted deck and, per slide, its largest embedded picture."""
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise UploadError("The presentation is corrupt or unreadable and could not be opened.") from exc
    slides = list(prs.slides)
    if not slides:
        raise UploadError("The presentation contains no slides.")
    if len(slides) > max_pages:
        raise UploadError(f"The presentation has {len(slides)} slides; the limit is {max_pages}.")

    slide_area = max(int(prs.slide_width or Emu(9144000)) * int(prs.slide_height or Emu(5143500)), 1)
    units: list[UnitContent] = []
    pictures: dict[int, bytes] = {}
    warnings: list[ExtractionWarning] = []
    ocr_missing: list[int] = []

    deadline = time.monotonic() + budget_seconds
    skipped: list[int] = []
    failed: list[int] = []
    for number, slide in enumerate(slides, start=1):
        allow_ocr = time.monotonic() < deadline
        try:
            unit, picture, no_ocr = _extract_slide(slide, number, slide_area, allow_ocr)
        except Exception:
            # One unreadable slide must not stop the deck: keep a placeholder and move on.
            log.warning("Slide %d could not be processed", number)
            unit, picture, no_ocr = UnitContent(number=number, warnings=["Slide could not be read."]), None, False
            failed.append(number)
        if picture is not None:
            pictures[number] = picture
        if no_ocr:
            ocr_missing.append(number)
        if not allow_ocr and unit.needs_visual:
            skipped.append(number)
        units.append(unit)
    if skipped:
        warnings.append(ExtractionWarning(
            severity="warning",
            message=f"The extraction time budget ({budget_seconds:.0f}s) ran out, so OCR was skipped for "
                    "these slides. Analysis continues with their text only.",
            units=skipped,
        ))
    if failed:
        warnings.append(ExtractionWarning(
            severity="warning",
            message="These slides could not be processed and are treated as empty. Analysis continues "
                    "with the rest of the presentation.",
            units=failed,
        ))

    if ocr_missing:
        warnings.append(ExtractionWarning(
            severity="warning",
            message="OCR is unavailable (install rapidocr-onnxruntime); image-based text on these "
                    "slides was not read.",
            units=sorted(set(ocr_missing)),
        ))
    deck = ExtractedDeck(
        filename=filename, source_type=source_type, unit_label="slide",  # type: ignore[arg-type]
        units=units, warnings=warnings,
    )
    return deck, pictures
